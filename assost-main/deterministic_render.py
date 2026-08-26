"""deterministic_render.py — التحويل من pages_flow الحتمي إلى HTML/PDF و DOCX.
الأحجام: width_frac من الأصل → عرض الصفحة فعلياً. صفحة الرسمة = صفحة كاملة.
"""
import os, sys, json, html, base64, re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

IMG_MARK = re.compile(r'\[IMG:([^|\]]+)(?:\|([^\]]+))?\]')


def _img_src(path):
    with open(path, 'rb') as f:
        return 'data:image/png;base64,' + base64.b64encode(f.read()).decode()


CSS = """
  @page { size: 410pt 600pt; margin: 34pt 30pt;
          @bottom-center { content: counter(page); font-family: 'Noto Naskh Arabic'; font-size: 9pt; } }
  @page cover { margin: 0; @bottom-center { content: none; } }
  .coverpage { page: cover; }
  body { font-family: 'Noto Naskh Arabic', 'Amiri', serif; font-size: 10.5pt;
         line-height: 1.65; text-align: justify; color: #1a1a1a; }
  h1.title { page: cover; text-align: center; font-size: 26pt; margin-top: 40%; }
  .illuspage, .textpage { page-break-before: always; }
  p { margin: 0 0 10pt; text-indent: 1.2em; text-align: justify;
      text-align-last: right; direction: rtl; }
  .illuspage { page: cover; page-break-after: always; }
  .illuspage img { display:block; margin:0 auto; }
  .illus-inline { text-align:center; margin: 10pt 0; page-break-inside: avoid; }
  .illus-inline img { max-width:100%; }
"""


def _marker_html(name, cls_attr, images_map):
    path = images_map.get(name) or os.path.join('images', name)
    if not os.path.exists(path):
        return f'<!-- MISSING:{name} -->'
    m = re.search(r'w([\d.]+)', cls_attr or '')
    frac = float(m.group(1)) if m else (1.0 if 'full_page' in (cls_attr or '') else 0.85)
    full = frac >= 0.98
    style = f'width:{frac*100:.0f}%;' + ('page-break-after:always;' if full else '')
    cls = 'illuspage' if full else 'illus-inline'
    return f'<div class="{cls}"><img style="{style}" src="{_img_src(path)}"></div>'


def flow_to_html(pages_flow, book_title='كتاب مترجم', cover_image=None,
                 chapter_titles=None):
    cover = ''
    if cover_image and os.path.exists(cover_image):
        cover = (f'<div class="coverpage"><img style="width:100%" '
                 f'src="{_img_src(cover_image)}"></div>')
    parts = [cover, f'<h1 class="title">{html.escape(book_title)}</h1>']
    for pg in pages_flow:
        images = pg.get('images', {})
        # كل وحدة صفحة أصلية تبدأ صفحة جديدة — شرط التطابق 1:1
        if pg.get('kind') == 'text':
            parts.append('<div style="page-break-before:always"></div>')
        for block in pg.get('text_with_markers', '').split('\n\n'):
            block = block.strip()
            if not block:
                continue
            m = IMG_MARK.fullmatch(block)
            if m:
                parts.append(_marker_html(m.group(1).strip(), m.group(2), images))
            else:
                parts.append(f'<p>{html.escape(block)}</p>')
    return f'<!DOCTYPE html><html dir="rtl" lang="ar"><head><meta charset="utf-8"><style>{CSS}</style></head><body>{"".join(parts)}</body></html>'


def render_pdf(pages_flow, output_path, book_title='كتاب مترجم', cover_image=None):
    from weasyprint import HTML
    html_doc = flow_to_html(pages_flow, book_title, cover_image)
    HTML(string=html_doc, base_url=os.path.dirname(output_path)).write_pdf(output_path)
    return os.path.getsize(output_path)


def render_docx(pages_flow, output_path, book_title='كتاب مترجم', cover_image=None):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    if cover_image and os.path.exists(cover_image):
        doc.add_picture(cover_image, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(book_title); r.font.size = Pt(28)
    doc.add_page_break()

    def rtl(p):
        pPr = p._p.get_or_add_pPr()
        b = pPr.makeelement('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi', {})
        b.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '1')
        pPr.append(b)

    for pi, pg in enumerate(pages_flow):
        for block in pg.get('text_with_markers', '').split('\n\n'):
            block = block.strip()
            if not block:
                continue
            m = IMG_MARK.fullmatch(block)
            if m:
                name, cls = m.group(1).strip(), m.group(2) or ''
                path = pg.get('images', {}).get(name) or os.path.join(
                    os.path.dirname(output_path), 'images', name)
                if os.path.exists(path):
                    mw = re.search(r'w([\d.]+)', cls)
                    frac = float(mw.group(1)) if mw else (
                        1.0 if 'full_page' in cls else 0.85)
                    doc.add_picture(path, width=Inches(min(5.8, 5.5 * max(frac, 0.3))))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; rtl(p)
            run = p.add_run(block); run.font.size = Pt(12); run.font.name = 'Amiri'
        if pi < len(pages_flow) - 1:
            doc.add_page_break()
    doc.save(output_path)
    return os.path.getsize(output_path)


if __name__ == '__main__':
    BASE = '/opt/data/projects/assost'
    sys.path.insert(0, '.')
    from deterministic_composer import build_pages_flow
    tr = json.load(open(f'{BASE}/tests/fullbook/paged_translations.json'))
    pf, rep = build_pages_flow(f'{BASE}/book.pdf', f'{BASE}/tests/fullbook/images', tr)
    out = f'{BASE}/tests/fullbook/deterministic'
    os.makedirs(out, exist_ok=True)
    szp = render_pdf(pf, f'{out}/det_book.pdf', 'Death March — المجلد الأول',
                     f'{BASE}/tests/fullbook/images/p1_x35.png')
    szd = render_docx(pf, f'{out}/det_book.docx', 'Death March — المجلد الأول',
                      f'{BASE}/tests/fullbook/images/p1_x35.png')
    print('PDF bytes:', szp, '| DOCX bytes:', szd)
